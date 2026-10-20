# -*- coding: utf-8 -*-
"""Generate Dice.com profile resume DOCX (one-off)."""
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt
from docx.oxml.ns import qn


def set_run_font(run, size_pt=11, bold=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size_pt)
    run.bold = bold


def add_heading_line(doc, text, size=14, bold=True, center=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
    run = p.add_run(text)
    set_run_font(run, size, bold)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_body(doc, text, bold=False, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, 11, bold)
    p.paragraph_format.space_after = Pt(space_after)
    return p


def add_bullet(doc, text):
    p = doc.add_paragraph(style="List Bullet")
    run = p.add_run(text)
    set_run_font(run, 11)
    p.paragraph_format.space_after = Pt(2)
    return p


def main():
    out = Path(__file__).resolve().parents[1] / "dice-profile-resume.docx"
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.85)
    section.right_margin = Inches(0.85)

    add_heading_line(doc, "YOUR FULL NAME", 16, True, center=True)
    add_heading_line(
        doc,
        "Full-Stack Software Engineer · SaaS · AI / Speech-to-Text",
        11,
        False,
        center=True,
    )
    contact = doc.add_paragraph()
    contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    parts = [
        "your.email@example.com",
        " | ",
        "City, State",
        " | ",
        "LinkedIn: linkedin.com/in/yourhandle",
        " | ",
        "Portfolio: ",
    ]
    for i, part in enumerate(parts):
        run = contact.add_run(part)
        set_run_font(run, 10, i == 0)
    link_run = contact.add_run("https://mp3totext.io")
    set_run_font(link_run, 10)
    link_run.underline = True
    contact.paragraph_format.space_after = Pt(14)

    add_heading_line(doc, "PROFESSIONAL SUMMARY", 12)
    add_body(
        doc,
        "Full-stack engineer with 7+ years building web products, APIs, and subscription "
        "SaaS. Recently focused on speech-to-text and media workflows: shipping a production "
        "transcription platform used by podcasters, researchers, and teams who need searchable "
        "transcripts, subtitles, and exports. Comfortable owning architecture, backend services, "
        "payments, and customer-facing UX in TypeScript / Node.js ecosystems.",
    )

    add_heading_line(doc, "TECHNICAL SKILLS", 12)
    add_bullet(
        doc,
        "Languages & runtime: TypeScript, JavaScript, SQL, Python (scripting); Node.js",
    )
    add_bullet(
        doc,
        "Frontend: React 19, Next.js (App Router), Tailwind CSS, accessible UI patterns",
    )
    add_bullet(
        doc,
        "Backend: REST APIs, auth sessions, webhooks, idempotent billing, background jobs",
    )
    add_bullet(
        doc,
        "Data: PostgreSQL, Drizzle ORM, Redis; schema design and migration discipline",
    )
    add_bullet(
        doc,
        "Cloud & media: object storage, async job pipelines, audio/video duration validation (FFprobe)",
    )
    add_bullet(
        doc,
        "Integrations: Stripe / payment providers, email, observability; monorepo (Turborepo-style workflows)",
    )
    add_bullet(doc, "Practices: i18n, SEO for marketing pages, admin tooling, security-minded API design")

    add_heading_line(doc, "PROFESSIONAL EXPERIENCE", 12)

    add_body(doc, "Founder & Lead Engineer — MP3ToText.io (mp3totext.io)", bold=True, space_after=2)
    add_body(doc, "Remote · 2024 – Present · https://mp3totext.io", bold=False, space_after=4)
    add_bullet(
        doc,
        "Built and operate MP3ToText.io, an online AI speech-to-text product: upload MP3/WAV/M4A/MP4 "
        "and related formats, receive timestamped transcripts with speaker labels, edit in-browser, "
        "and export TXT, PDF, DOCX, CSV, SRT, and VTT.",
    )
    add_bullet(
        doc,
        "Designed a freemium model (guest trials, free accounts, minute-based subscriptions and one-time "
        "packs) with credit accounting, renewal logic, and admin visibility into users, orders, and jobs.",
    )
    add_bullet(
        doc,
        "Implemented a split Next.js web + API monorepo: marketing/SEO landing pages, authenticated dashboard, "
        "transcription job lifecycle, streaming playback, and internal admin for support and operations.",
    )
    add_bullet(
        doc,
        "Integrated payment webhooks, subscription state machines (active / past_due / cancel-at-period-end), "
        "and scheduled jobs for yearly-plan monthly credit grants.",
    )
    add_bullet(
        doc,
        "Shipped product extensions aligned with SEO and user demand (e.g., MP3-to-SRT and M4A-to-text flows) "
        "while keeping shared transcription infrastructure.",
    )
    add_bullet(
        doc,
        "Stack highlights: TypeScript, Next.js, PostgreSQL, Drizzle, Redis, Better Auth, Trigger.dev-style "
        "schedulers, R2-compatible object storage.",
    )

    add_body(doc, "Senior Software Engineer — B2B SaaS (Product & Platform team)", bold=True, space_after=2)
    add_body(doc, "Remote / Hybrid · 2020 – 2024", bold=False, space_after=4)
    add_bullet(
        doc,
        "Delivered customer-facing features and internal admin tools for a multi-tenant SaaS product "
        "(role-based access, usage metering, reporting).",
    )
    add_bullet(
        doc,
        "Owned API design and PostgreSQL performance for high-read dashboards; reduced p95 latency on "
        "list endpoints through indexing and query refactors.",
    )
    add_bullet(
        doc,
        "Partnered with design and support on onboarding flows, error messaging, and release checklists; "
        " mentored two mid-level engineers on code review and testing habits.",
    )

    add_body(doc, "Software Engineer — Web Applications", bold=True, space_after=2)
    add_body(doc, "2017 – 2020", bold=False, space_after=4)
    add_bullet(
        doc,
        "Built React/Node features for content and collaboration tools: file uploads, notifications, "
        "and third-party OAuth integrations.",
    )
    add_bullet(
        doc,
        "Maintained CI pipelines and incremental TypeScript adoption across a legacy JavaScript codebase.",
    )

    add_heading_line(doc, "SELECTED PROJECT (PUBLIC)", 12)
    add_body(doc, "MP3ToText.io — https://mp3totext.io", bold=True, space_after=2)
    add_body(
        doc,
        "Public portfolio reference for current work: AI transcription SaaS with 150+ language support, "
        "subtitle workflows, summaries/mind maps, and enterprise-friendly export options. "
        "Contact for the product: support@mp3totext.io.",
        space_after=8,
    )

    add_heading_line(doc, "EDUCATION", 12)
    add_body(doc, "B.S., Computer Science (or related field) — Your University Name, Graduation Year")

    add_heading_line(doc, "ADDITIONAL", 12)
    add_bullet(doc, "Authorized to work in [country]; open to remote-friendly roles in full-stack / platform engineering.")
    add_bullet(doc, "Dice profile and resume highlight ongoing indie SaaS ownership alongside prior team experience.")

    doc.save(out)
    print(f"Wrote {out}")


if __name__ == "__main__":
    main()
